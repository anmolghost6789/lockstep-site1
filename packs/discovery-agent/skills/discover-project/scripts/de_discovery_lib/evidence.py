"""Content-addressed evidence ingestion and bounded text extraction."""

from __future__ import annotations

import hashlib
import mimetypes
import os
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from .io import DiscoveryError, RuntimePaths, atomic_write_text, utc_now
from .security import scan_text
from .state import save_state


TEXT_SUFFIXES = {
    ".csv",
    ".html",
    ".ini",
    ".json",
    ".log",
    ".md",
    ".py",
    ".rst",
    ".sql",
    ".toml",
    ".tsv",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}
IGNORED_DIR_NAMES = {
    ".git",
    ".de-discovery",
    ".idea",
    ".vscode",
    "__pycache__",
    "node_modules",
}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DiscoveryError("Text encoding is not supported")


def _xml_text(xml_bytes: bytes) -> str:
    root = ElementTree.fromstring(xml_bytes)
    values = [node.text for node in root.iter() if node.text and node.text.strip()]
    return "\n".join(values)


def _natural_key(value: str) -> list[str | int]:
    return [int(part) if part.isdigit() else part for part in re.split(r"(\d+)", value)]


def _validate_archive_size(path: Path, max_uncompressed_bytes: int) -> None:
    with zipfile.ZipFile(path) as archive:
        total = sum(item.file_size for item in archive.infolist())
    if total > max_uncompressed_bytes:
        raise DiscoveryError(
            f"Archive expands beyond safety limit {max_uncompressed_bytes} bytes: {path}"
        )


def _extract_docx(path: Path, max_uncompressed_bytes: int) -> str:
    _validate_archive_size(path, max_uncompressed_bytes)
    parts: list[str] = []
    with zipfile.ZipFile(path) as archive:
        names = [
            name
            for name in archive.namelist()
            if name == "word/document.xml"
            or re.match(r"word/(?:header|footer|comments)\d*\.xml$", name)
        ]
        for name in sorted(names, key=_natural_key):
            parts.append(f"\n## {name}\n")
            parts.append(_xml_text(archive.read(name)))
    return "\n".join(parts)


def _extract_pptx(path: Path, max_uncompressed_bytes: int) -> str:
    _validate_archive_size(path, max_uncompressed_bytes)
    parts: list[str] = []
    with zipfile.ZipFile(path) as archive:
        names = [
            name
            for name in archive.namelist()
            if re.match(r"ppt/(?:slides/slide|notesSlides/notesSlide)\d+\.xml$", name)
        ]
        for name in sorted(names, key=_natural_key):
            parts.append(f"\n## {name}\n")
            parts.append(_xml_text(archive.read(name)))
    return "\n".join(parts)


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DiscoveryError(
            "PDF extraction requires pypdf; install scripts/requirements.txt"
        ) from exc
    reader = PdfReader(str(path))
    return "\n\n".join(
        f"## Page {number}\n{page.extract_text() or ''}"
        for number, page in enumerate(reader.pages, start=1)
    )


def _extract_xlsx(path: Path, max_rows: int, max_uncompressed_bytes: int) -> str:
    _validate_archive_size(path, max_uncompressed_bytes)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DiscoveryError(
            "XLSX extraction requires openpyxl; install scripts/requirements.txt"
        ) from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    remaining = max_rows
    try:
        for worksheet in workbook.worksheets:
            parts.append(f"\n## Sheet: {worksheet.title}\n")
            for row in worksheet.iter_rows(values_only=True):
                if remaining <= 0:
                    parts.append("\n[Extraction row limit reached]\n")
                    return "\n".join(parts)
                parts.append("\t".join("" if value is None else str(value) for value in row[:100]))
                remaining -= 1
    finally:
        workbook.close()
    return "\n".join(parts)


def extract_text(path: Path, max_rows: int, max_chars: int) -> str:
    suffix = path.suffix.lower()
    max_archive_bytes = max(10_000_000, max_chars * 8)
    if suffix in TEXT_SUFFIXES:
        return _decode_text(path.read_bytes())
    if suffix == ".docx":
        return _extract_docx(path, max_archive_bytes)
    if suffix == ".pptx":
        return _extract_pptx(path, max_archive_bytes)
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _extract_xlsx(path, max_rows, max_archive_bytes)
    raise DiscoveryError(f"Unsupported evidence format: {suffix or '[no extension]'}")


def _safe_filename(name: str) -> str:
    clean = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("._")
    return clean[:180] or "source"


def _logical_source_key(path: Path) -> str:
    return hashlib.sha256(str(path.resolve()).casefold().encode("utf-8")).hexdigest()[:24]


def _verified_snapshot(source: Path, destination: Path, digest: str) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        if _sha256(destination) != digest:
            raise DiscoveryError(f"Existing content-addressed evidence is corrupt: {destination}")
        return
    fd, temp_name = tempfile.mkstemp(
        prefix=f".{destination.name}.",
        suffix=".tmp",
        dir=destination.parent,
    )
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        shutil.copy2(source, temp_path)
        if _sha256(temp_path) != digest:
            raise DiscoveryError(f"Source changed while it was being captured: {source}")
        os.replace(temp_path, destination)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def discover_files(
    input_path: Path,
    *,
    project_root: Path,
    knowledge_root: Path,
    max_files: int,
) -> list[Path]:
    resolved = input_path.expanduser().resolve()
    if not resolved.exists():
        raise DiscoveryError(f"Input path does not exist: {resolved}")
    excluded_roots = {knowledge_root.resolve(), (project_root / ".de-discovery").resolve()}
    if resolved.is_file():
        if any(root == resolved or root in resolved.parents for root in excluded_roots):
            raise DiscoveryError(f"Refusing to ingest discovery-owned output: {resolved}")
        return [resolved]

    files: list[Path] = []
    for candidate in sorted(resolved.rglob("*")):
        if not candidate.is_file():
            continue
        if any(part in IGNORED_DIR_NAMES for part in candidate.parts):
            continue
        if any(root == candidate or root in candidate.parents for root in excluded_roots):
            continue
        files.append(candidate)
        if len(files) > max_files:
            raise DiscoveryError(f"Input exceeds configured max_files={max_files}")
    return files


def ingest_file(
    source_path: Path,
    *,
    paths: RuntimePaths,
    config: dict[str, Any],
    state: dict[str, Any],
) -> dict[str, Any]:
    source = source_path.expanduser().resolve()
    inputs = config["inputs"]
    size = source.stat().st_size
    if size > inputs["max_file_bytes"]:
        raise DiscoveryError(
            f"Evidence file exceeds max_file_bytes={inputs['max_file_bytes']}: {source}"
        )

    digest = _sha256(source)
    source_key = _logical_source_key(source)
    registry = state["sources"].get(source_key)
    if registry and registry.get("current_revision"):
        current = next(
            (
                item
                for item in registry["revisions"]
                if item.get("id") == registry["current_revision"]
            ),
            None,
        )
        if current and current.get("sha256") == digest:
            return {
                "status": "unchanged",
                "path": str(source),
                "revision_id": current["id"],
                "sha256": digest,
            }

    extracted = extract_text(
        source,
        inputs["max_tabular_rows"],
        inputs["max_extracted_chars"],
    )
    findings = scan_text(extracted)
    if findings:
        locations = ", ".join(f"line {item.line} ({item.detector})" for item in findings[:10])
        raise DiscoveryError(
            f"Potential secret detected in {source}; ingestion refused at {locations}. "
            "Redact the source and retry."
        )
    if _sha256(source) != digest:
        raise DiscoveryError(f"Source changed while it was being extracted: {source}")
    if len(extracted) > inputs["max_extracted_chars"]:
        extracted = extracted[: inputs["max_extracted_chars"]] + "\n\n[Extraction character limit reached]\n"

    registry = state["sources"].setdefault(
        source_key,
        {"logical_path": str(source), "current_revision": None, "revisions": []},
    )
    previous = registry.get("current_revision")
    revision_id = f"source-{source_key[:8]}-{digest[:16]}"
    filename = _safe_filename(source.name)
    mode = inputs["evidence_mode"]
    if mode == "snapshot":
        snapshot = paths.evidence_root / "sha256" / digest / filename
        _verified_snapshot(source, snapshot, digest)
        resource = f"evidence://sha256/{digest}/{filename}"
    else:
        snapshot = None
        resource = source.as_uri().replace("file:", "file+sha256:", 1) + f"?digest={digest}"

    derived = paths.evidence_root / "derived" / f"{digest}.md"
    if not derived.exists():
        atomic_write_text(
            derived,
            (
                f"# Extracted evidence: {filename}\n\n"
                f"- Source revision: `{revision_id}`\n"
                f"- SHA-256: `{digest}`\n"
                f"- Resource: `{resource}`\n\n"
                "## Extracted content\n\n"
                f"{extracted}\n"
            ),
        )

    revision = {
        "id": revision_id,
        "sha256": digest,
        "captured_at": utc_now(),
        "original_path": str(source),
        "resource": resource,
        "snapshot_path": str(snapshot) if snapshot else None,
        "derived_path": str(derived),
        "size_bytes": size,
        "media_type": mimetypes.guess_type(source.name)[0] or "application/octet-stream",
    }
    if not any(item.get("id") == revision_id for item in registry["revisions"]):
        registry["revisions"].append(revision)
    registry["current_revision"] = revision_id
    if previous:
        superseded = next(
            (
                item
                for item in state.get("pending_source_revisions", [])
                if item.get("revision_id") == previous
            ),
            None,
        )
        if superseded:
            state["source_reviews"].append(
                {
                    **superseded,
                    "disposition": "superseded-before-review",
                    "concepts": [],
                    "reason": f"Superseded by {revision_id} before semantic review",
                    "reviewed_at": utc_now(),
                }
            )
    state["pending_source_revisions"] = [
        item
        for item in state.get("pending_source_revisions", [])
        if item.get("revision_id") not in {revision_id, previous}
    ]
    state["pending_source_revisions"].append(
        {
            "revision_id": revision_id,
            "previous_revision_id": previous,
            "logical_path": str(source),
            "status": "changed" if previous else "new",
            "registered_at": utc_now(),
        }
    )
    if state.get("phase") == "complete":
        state["phase"] = "discover"
        state["last_completed_step"] = "New evidence requires review"
    save_state(paths.state_path, state)
    return {
        "status": "changed" if previous else "new",
        "path": str(source),
        "previous_revision_id": previous,
        "revision_id": revision_id,
        "sha256": digest,
        "resource": resource,
        "derived_path": str(derived),
    }
