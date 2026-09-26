#!/usr/bin/env python3
"""Canonical workbook generator for the design-artifact phase.

Consumes the canonical design-phase JSONs that /design-architecture produces:

* ``outputs/00_state/design/column_mappings.json``  — ``{"layers": [...],
  "tables": [...], "mappings": [entry, ...]}``. Per-layer metadata (title,
  schema, database) lives in ``layers``; per-table metadata (purpose, grain,
  key strategy, load strategy, frequency, schema, database, table type) lives
  in ``tables``; one entry per column lives in ``mappings``.
* ``outputs/00_state/design/dq_rules_design.json``  — ``{"rules": [entry, ...]}``

``target_model_design.md`` is a RENDERED VIEW (see ``render_design_md.py``)
and is never read by this generator. The generator degrades gracefully to
``-`` placeholders when a metadata block or field is absent.

Produces per layer: ``STTM_{layer}.xlsx``, ``DATA_MODEL_{layer}.xlsx``,
``DQ_{layer}.xlsx``, ``ER_DIAGRAM_{layer}.md``; and at the artifacts root:
``ER_DIAGRAM.xlsx`` (all layers) and ``ER_DIAGRAM_LINEAGE.md`` (all layers).

Invariants:
* Templates are loaded via ``load_workbook`` and their sheets are copied —
  all formatting (theme, palette, widths) comes from the template, never
  from hardcoded fills in this script.
* ``_template_reference`` is deleted before save; final workbooks contain
  only ``table_tracker`` plus one sheet per current-run table.
* DQ Rule Expression cells hold BARE boolean predicates (design-architecture rule R8).
  The verifier / DQ runtime wraps them (``SELECT ... WHERE (filter) AND
  (predicate)``); the generator never writes a leading ``WHERE`` into a cell.
* ``*_reason`` design rationale never becomes a workbook column; for
  NO_DQ_APPLICABLE / [NEEDS_HUMAN_REVIEW] rules the reason goes into the
  existing Comments column.
* All text/JSON IO is explicit ``encoding='utf-8'``.
* Deterministic: same inputs produce the same workbook contents.

Usage:
    python generate_workbooks.py --layer L0 \
        --design-dir outputs/00_state/design \
        --out-root outputs/05_artifacts \
        --templates-dir .claude/skills/design-agent/templates/workbooks

    ``--layer all`` generates every layer found in column_mappings.json.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    print(
        "generate_workbooks: missing required package 'openpyxl'. Install with:\n"
        "  pip install -r .claude/skills/design-agent/scripts/requirements.txt",
        file=sys.stderr,
    )
    sys.exit(2)

DQ_CHECK_LABEL = {
    "NOT_NULL": "Not NULL", "UNIQUENESS": "Uniqueness", "FK_INTEGRITY": "Referential Integrity",
    "RANGE": "Range Check", "REGEX_MATCH": "Pattern Check", "VALUE_IN_SET": "Value In Set",
    "CONSISTENCY": "Consistency Check", "CUSTOM_SQL": "Custom SQL",
    "ROW_COUNT_NOT_ZERO": "Row Count Check", "FRESHNESS": "Freshness Check",
    "COMPLETENESS_%": "Completeness Check", "DATE_FORMAT": "Date Format Check",
    "NO_DQ_APPLICABLE": "NO_DQ_APPLICABLE",
}
CRIT_LABEL = {"CRITICAL": "C", "NON_CRITICAL": "NC", "C": "C", "NC": "NC", "INFO": "INFO"}
ACRONYMS = {"id": "ID", "sk": "SK", "fk": "FK", "dq": "DQ", "qc": "QC", "kpi": "KPI",
            "md5": "MD5", "loa": "LOA", "ym": "YM", "atc": "ATC"}
TABLE_TYPE_PREFIXES = (
    ("d_", "dimension"), ("f_", "fact"), ("ref_", "reference"),
    ("rw_", "raw"), ("raw_", "raw"), ("stg_", "staging"), ("cnf_", "conformed"),
)


def _read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _yn(value) -> str:
    """Normalize a nullable/pk flag (bool, Y/N, true/false strings) to Y/N."""
    if isinstance(value, bool):
        return "Y" if value else "N"
    text = str(value or "").strip().upper()
    return "Y" if text in {"Y", "YES", "TRUE", "1"} else "N"


def _dash(value) -> str:
    text = str(value if value is not None else "").strip()
    return text if text else "-"


def logical_name(col: str) -> str:
    return " ".join(ACRONYMS.get(p, p.capitalize()) for p in col.split("_"))


def role_of(col: dict) -> str:
    if col.get("role"):
        return col["role"]
    name = col["name"].lower()
    if name.startswith("sys_") or name.startswith("audit_"):
        return "audit"
    if col["pk"] == "Y":
        return "surrogate_key" if name.endswith("_sk") or col["data_type"].upper() == "UUID" else "business_key"
    return "attribute"


def table_type_of(table_name: str, explicit: str) -> str:
    if explicit and explicit != "-":
        return explicit
    lname = table_name.lower()
    for prefix, ttype in TABLE_TYPE_PREFIXES:
        if lname.startswith(prefix):
            return ttype
    return "-"


def bare_predicate(expression: str) -> str:
    """Enforce design-architecture rule R8: strip a leading WHERE so the cell is a bare predicate."""
    expr = (expression or "").strip()
    if expr.upper().startswith("WHERE "):
        expr = expr[6:].strip()
    return expr


# ---------------------------------------------------------------------------
# Design assembly from column_mappings.json + dq_rules_design.json
# ---------------------------------------------------------------------------

def load_design(design_dir: Path) -> dict:
    """Build the in-memory design model {layers: {id: {title, schema, database, tables}}}."""
    mapping_path = design_dir / "column_mappings.json"
    if not mapping_path.exists():
        raise SystemExit(f"FATAL: missing {mapping_path}")
    doc = _read_json(mapping_path)
    mappings = doc.get("mappings") or []
    if not mappings:
        raise SystemExit(f"FATAL: {mapping_path} has no 'mappings' entries")

    rules_path = design_dir / "dq_rules_design.json"
    rules = _read_json(rules_path).get("rules") or [] if rules_path.exists() else []
    if not rules_path.exists():
        print(f"WARNING: {rules_path} missing — DQ workbooks will be empty", file=sys.stderr)

    # Per-layer / per-table metadata blocks (design-architecture JSON-as-source contract).
    layer_meta: dict[str, dict] = {}
    for lmeta in doc.get("layers") or []:
        lid = str(lmeta.get("id") or "").strip()
        if lid:
            layer_meta[lid] = lmeta
    table_meta: dict[tuple[str, str], dict] = {}
    for tmeta in doc.get("tables") or []:
        lid = str(tmeta.get("layer") or "").strip()
        tname = str(tmeta.get("table") or tmeta.get("target_table") or "").strip()
        if lid and tname:
            table_meta[(lid, tname)] = tmeta

    layers: dict[str, dict] = {}
    for entry in mappings:
        layer_id = str(entry.get("layer") or "").strip()
        table = str(entry.get("target_table") or "").strip()
        column = str(entry.get("target_column") or "").strip()
        if not layer_id or not table or not column:
            print(f"WARNING: skipping mapping entry missing layer/table/column: {entry}", file=sys.stderr)
            continue
        lmeta = layer_meta.get(layer_id, {})
        layer = layers.setdefault(layer_id, {
            "title": str(lmeta.get("title") or "").strip() or f"Layer {layer_id}",
            "schema": str(lmeta.get("schema") or "").strip() or "-",
            "database": str(lmeta.get("database") or "").strip() or "-",
            "tables": {},
        })
        tmeta = table_meta.get((layer_id, table), {})
        t = layer["tables"].setdefault(table, {
            "purpose": _dash(tmeta.get("purpose")),
            "load_strategy": _dash(tmeta.get("load_strategy")),
            "frequency": _dash(tmeta.get("frequency")),
            "table_type": table_type_of(table, _dash(tmeta.get("table_type"))),
            "schema": _dash(tmeta.get("schema")) if _dash(tmeta.get("schema")) != "-" else layer["schema"],
            "database": _dash(tmeta.get("database")) if _dash(tmeta.get("database")) != "-" else layer["database"],
            "columns": [],
            "rules": [],
        })
        t["columns"].append({
            "name": column,
            "data_type": _dash(entry.get("data_type")),
            "nullable": _yn(entry.get("nullable")),
            "pk": _yn(entry.get("is_primary_key")),
            "role": str(entry.get("role") or "").strip(),
            "source_system": _dash(entry.get("source_system")),
            "source_table": _dash(entry.get("source_table")),
            "source_column": _dash(entry.get("source_column")),
            "source_data_type": _dash(entry.get("source_data_type")),
            "transformation": _dash(entry.get("transformation")),
            "filter_conditions": _dash(entry.get("filter_conditions")),
            "join_conditions": _dash(entry.get("join_conditions")),
            "source_reference": _dash(entry.get("source_reference")),
            "description": str(entry.get("description") or "").strip(),
            "notes": str(entry.get("notes") or "").strip(),
        })

    for rule in rules:
        layer_id = str(rule.get("layer") or "").strip()
        table = str(rule.get("target_table") or "").strip()
        if layer_id not in layers or table not in layers[layer_id].get("tables", {}):
            print(
                f"WARNING: DQ rule {rule.get('rule_id')} targets unknown table "
                f"{layer_id}.{table} — skipped (fix column_mappings.json / dq_rules_design.json)",
                file=sys.stderr,
            )
            continue
        rule_type = str(rule.get("rule_type") or "").strip()
        raw_expr = str(rule.get("dq_rule_expression") or "").strip()
        if rule_type == "CUSTOM_SQL" and raw_expr.upper().startswith("[CUSTOM_SQL"):
            expression = str(rule.get("custom_sql") or "").strip()
            if not expression:
                expression = "[NEEDS_HUMAN_REVIEW]"
        elif raw_expr.startswith("["):
            # [NEEDS_HUMAN_REVIEW] / [NO_DQ_APPLICABLE] markers pass through verbatim
            expression = raw_expr
        else:
            expression = bare_predicate(raw_expr)
        comments: list[str] = []
        rule_id = str(rule.get("rule_id") or "").strip()
        if rule_id:
            comments.append(rule_id)
        for reason_field in ("dq_rule_expression_reason", "threshold_value_reason"):
            reason = str(rule.get(reason_field) or "").strip()
            if reason and reason not in {"-", "not applicable; NO_DQ_APPLICABLE marker"}:
                comments.append(reason)
        if rule_type == "UNIQUENESS":
            comments.append(f"Evaluated as: GROUP BY {rule.get('target_column')} HAVING COUNT(*) > 1")
        layers[layer_id]["tables"][table]["rules"].append({
            "rule_id": rule_id,
            "column": _dash(rule.get("target_column")),
            "check_type": rule_type,
            "expression": expression,
            "severity": str(rule.get("severity") or "").strip(),
            "threshold": _dash(rule.get("threshold_value")),
            "reference": _dash(rule.get("source_reference")),
            "comments": " | ".join(comments),
        })

    for layer_id, layer in layers.items():
        for table, t in layer["tables"].items():
            if not t["rules"]:
                print(f"WARNING: {layer_id}.{table} has no DQ rules (design-architecture rule R6 floor)", file=sys.stderr)

    return {"layers": layers}


def relationships(design: dict) -> list[dict]:
    """Derive (from_table, from_col, to_table, to_col, kind, ref) from Join Conditions cells."""
    rel_re = re.compile(r"([A-Za-z_][\w]*)\.([\w]+)\s*(?:=|->)\s*([A-Za-z_][\w]*)\.([\w]+)")
    short_rel_re = re.compile(r"(\w+)\s*->\s*([A-Za-z_][\w]*)\.([\w]+)")
    rels, seen = [], set()
    for lname, layer in design["layers"].items():
        for tname, t in layer["tables"].items():
            for col in t["columns"]:
                jc = col["join_conditions"]
                if not jc or jc == "-":
                    continue
                kind = "FK" if jc.startswith("FK:") else "semantic"
                matched = False
                for g in rel_re.finditer(jc):
                    a_t, a_c, b_t, b_c = g.groups()
                    ft, fc, tt, tc = a_t, a_c, b_t, b_c
                    if ft != tname and tt == tname:
                        ft, fc, tt, tc = b_t, b_c, a_t, a_c
                    if ft == tt:
                        continue
                    matched = True
                    key = (ft, fc, tt, tc)
                    if key not in seen:
                        seen.add(key)
                        rels.append({"from_table": ft, "from_column": fc, "to_table": tt,
                                     "to_column": tc, "kind": kind, "layer": lname,
                                     "reference": col["source_reference"]})
                if not matched:
                    g2 = short_rel_re.search(jc)
                    if g2:
                        key = (tname, g2.group(1), g2.group(2), g2.group(3))
                        if key not in seen:
                            seen.add(key)
                            rels.append({"from_table": tname, "from_column": g2.group(1),
                                         "to_table": g2.group(2), "to_column": g2.group(3),
                                         "kind": "semantic", "layer": lname,
                                         "reference": col["source_reference"]})
    return rels


# ---------------------------------------------------------------------------
# Workbook writers (template-clone based; theme comes from the template)
# ---------------------------------------------------------------------------

def clear_rows(ws, first: int) -> None:
    if ws.max_row >= first:
        ws.delete_rows(first, ws.max_row - first + 1)


def copy_ref(wb, name: str):
    sheet_name = name[:31]
    if sheet_name in wb.sheetnames:
        raise SystemExit(f"FATAL: duplicate sheet name after truncation: {sheet_name}")
    ws = wb.copy_worksheet(wb["_template_reference"])
    ws.title = sheet_name
    return ws


def tracker(wb, rows: list[tuple]) -> None:
    ws = wb["table_tracker"]
    clear_rows(ws, 2)
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)


def finalize(wb, out: Path, planned: set[str]) -> None:
    if "_template_reference" in wb.sheetnames:
        del wb["_template_reference"]
    bad = [s for s in wb.sheetnames if s.startswith("_") or s not in planned]
    if bad:
        raise SystemExit(f"FATAL: unexpected sheets {bad} for {out}")
    wb.save(out)


def gen_sttm(layer_id: str, layer: dict, tpl_dir: Path, out_dir: Path) -> None:
    wb = load_workbook(tpl_dir / "STTM_TEMPLATE.xlsx")
    for tname, t in layer["tables"].items():
        ws = copy_ref(wb, tname)
        ws["C2"], ws["C3"], ws["C4"] = layer_id, t["purpose"], t["load_strategy"]
        ws["C5"], ws["C6"], ws["C7"] = t["schema"], t["database"], t["frequency"]
        clear_rows(ws, 10)
        for r, c in enumerate(t["columns"], start=10):
            vals = [tname, c["name"], role_of(c), c["data_type"], c["nullable"],
                    c["source_table"], c["source_column"], c["source_data_type"],
                    c["transformation"], c["filter_conditions"], c["join_conditions"],
                    c["description"], c["pk"], t["schema"], t["database"],
                    c["source_reference"], c["notes"]]
            for j, v in enumerate(vals, start=1):
                ws.cell(row=r, column=j, value=v)
    tracker(wb, [(i, n, b["table_type"]) for i, (n, b) in enumerate(layer["tables"].items(), 1)])
    finalize(wb, out_dir / f"STTM_{layer_id}.xlsx", {"table_tracker", *[n[:31] for n in layer["tables"]]})


def gen_dm(layer_id: str, layer: dict, tpl_dir: Path, out_dir: Path) -> None:
    wb = load_workbook(tpl_dir / "DATA_MODEL_TEMPLATE.xlsx")
    for tname, t in layer["tables"].items():
        ws = copy_ref(wb, tname)
        ws["C2"], ws["C3"], ws["C4"] = layer_id, t["purpose"], t["load_strategy"]
        ws["C5"], ws["C6"], ws["C7"] = t["schema"], t["database"], t["frequency"]
        clear_rows(ws, 10)
        for n, c in enumerate(t["columns"], start=1):
            vals = [n, t["schema"], t["database"], tname, logical_name(c["name"]),
                    c["name"], c["data_type"], role_of(c), c["nullable"], c["pk"],
                    c["description"], n, c["source_reference"], c["notes"]]
            for j, v in enumerate(vals, start=1):
                ws.cell(row=9 + n, column=j, value=v)
    tracker(wb, [(i, n, b["table_type"]) for i, (n, b) in enumerate(layer["tables"].items(), 1)])
    finalize(wb, out_dir / f"DATA_MODEL_{layer_id}.xlsx", {"table_tracker", *[n[:31] for n in layer["tables"]]})


def gen_dq(layer_id: str, layer: dict, tpl_dir: Path, out_dir: Path) -> None:
    wb = load_workbook(tpl_dir / "DQ_TEMPLATE.xlsx")
    for tname, t in layer["tables"].items():
        dtype = {c["name"]: c["data_type"] for c in t["columns"]}
        ws = copy_ref(wb, tname)
        ws["C2"], ws["C3"], ws["C4"] = layer_id, tname, t["purpose"]
        ws["C5"], ws["C6"] = t["schema"], t["database"]
        clear_rows(ws, 9)
        for n, rule in enumerate(t["rules"], start=1):
            vals = [n, t["schema"], t["database"], tname, rule["column"],
                    dtype.get(rule["column"], "-"),
                    DQ_CHECK_LABEL.get(rule["check_type"], rule["check_type"]),
                    rule["expression"],
                    CRIT_LABEL.get(rule["severity"], rule["severity"]),
                    rule["threshold"], rule["reference"], rule["comments"]]
            for j, v in enumerate(vals, start=1):
                ws.cell(row=8 + n, column=j, value=v)
    tracker(wb, [(i, n, f"{len(b['rules'])} rules") for i, (n, b) in enumerate(layer["tables"].items(), 1)])
    finalize(wb, out_dir / f"DQ_{layer_id}.xlsx", {"table_tracker", *[n[:31] for n in layer["tables"]]})


def sanitize_type(dt: str) -> str:
    return re.sub(r"[^A-Za-z0-9_]", "", dt.replace("(", "_").replace(",", "_").replace(")", ""))


def gen_er_xlsx(design: dict, rels: list[dict], tpl_dir: Path, out_root: Path) -> None:
    """ER_DIAGRAM.xlsx covers ALL layers (one row per table across all layers)."""
    wb = load_workbook(tpl_dir / "ER_DIAGRAM_TEMPLATE.xlsx")
    fk_cols = {(r["from_table"], r["from_column"]): (r["to_table"], r["to_column"]) for r in rels}
    rel_layer = {}
    for r in rels:
        rel_layer[(r["from_table"], r["from_column"])] = r["layer"]

    lo = wb["lineage_overview"]
    clear_rows(lo, 11)
    row = 11
    for layer_id, layer in design["layers"].items():
        for tname, t in layer["tables"].items():
            srcs = sorted({c["source_table"] for c in t["columns"] if c["source_table"] not in ("-", "")})
            keys = [c["name"] for c in t["columns"] if c["pk"] == "Y"]
            vals = [layer_id, tname, t["table_type"], "; ".join(srcs) or "- (system generated)",
                    ", ".join(keys), len(t["columns"]), t["purpose"], "design/column_mappings.json"]
            for j, v in enumerate(vals, start=1):
                lo.cell(row=row, column=j, value=v)
            row += 1

    tr = wb["table_relationships"]
    clear_rows(tr, 11)
    for i, r in enumerate(rels):
        vals = [r["layer"], r["from_table"], r["from_column"],
                "FK" if r["kind"] == "FK" else "semantic FK",
                r["layer"], r["to_table"], r["to_column"], r["reference"]]
        for j, v in enumerate(vals, start=1):
            tr.cell(row=11 + i, column=j, value=v)

    ld = wb["layer_diagram_data"]
    clear_rows(ld, 11)
    row = 11
    for layer_id, layer in design["layers"].items():
        for tname, t in layer["tables"].items():
            for c in t["columns"]:
                fk = fk_cols.get((tname, c["name"]))
                vals = [tname, layer_id, c["name"], c["data_type"], c["pk"],
                        "Y" if fk else "N", f"{fk[0]}.{fk[1]}" if fk else "", c["source_reference"]]
                for j, v in enumerate(vals, start=1):
                    ld.cell(row=row, column=j, value=v)
                row += 1

    finalize(wb, out_root / "ER_DIAGRAM.xlsx",
             {"table_tracker", "lineage_overview", "table_relationships", "layer_diagram_data"})


def gen_er_md(layer_id: str, layer: dict, rels: list[dict], out_dir: Path) -> None:
    fk_cols = {(r["from_table"], r["from_column"]) for r in rels}
    lines = [f"# ER Diagram - {layer_id}", "", "```mermaid", "erDiagram"]
    for tname, t in layer["tables"].items():
        lines.append(f"  {tname} {{")
        for c in t["columns"]:
            marks = [m for m, ok in (("PK", c["pk"] == "Y"), ("FK", (tname, c["name"]) in fk_cols)) if ok]
            suffix = " " + ",".join(marks) if marks else ""
            lines.append(f"    {sanitize_type(c['data_type'])} {c['name']}{suffix}")
        lines.append("  }")
    for r in rels:
        label = "references" if r["kind"] == "FK" else "semantically_references"
        lines.append(f'  {r["from_table"]} }}o--|| {r["to_table"]} : "{r["from_column"]} {label} {r["to_column"]}"')
    inferred = [f"- {r['from_table']}.{r['from_column']} -> {r['to_table']}.{r['to_column']} is INFERRED"
                for r in rels if "INFERRED" in r["reference"].upper()]
    lines += ["```", ""]
    if inferred:
        lines += ["## Known gaps"] + inferred + [""]
    (out_dir / f"ER_DIAGRAM_{layer_id}.md").write_text("\n".join(lines), encoding="utf-8")


def gen_lineage_md(design: dict, out_root: Path) -> None:
    def nid(s: str) -> str:
        return re.sub(r"[^A-Za-z0-9_]", "_", s)

    lines = ["# ER Diagram - Whole-Flow Lineage", "", "```mermaid", "flowchart LR"]
    layer_of_table = {t: lname for lname, layer in design["layers"].items() for t in layer["tables"]}
    ext_srcs: dict[str, set[str]] = {}
    for lname, layer in design["layers"].items():
        for tname, t in layer["tables"].items():
            for c in t["columns"]:
                st = c["source_table"]
                if (st not in ("-", "") and st not in layer_of_table
                        and not c["transformation"].startswith(("System generated", "Manually inserted"))):
                    ext_srcs.setdefault(st, set()).add(tname)
    if ext_srcs:
        lines.append('  subgraph SRC["Source Systems"]')
        for s in sorted(ext_srcs):
            lines.append(f'    {nid(s)}["{s}"]')
        lines.append("  end")
    for lname, layer in design["layers"].items():
        lines.append(f'  subgraph {nid(lname)}["{layer["title"]}"]')
        for tname in layer["tables"]:
            lines.append(f'    {nid(tname)}["{tname}"]')
        lines.append("  end")
    # cross-layer lineage edges: source table -> target table (short labels)
    for lname, layer in design["layers"].items():
        for tname, t in layer["tables"].items():
            upstream = sorted({c["source_table"] for c in t["columns"]
                               if c["source_table"] in layer_of_table and c["source_table"] != tname})
            for u in upstream:
                lines.append(f"  {nid(u)} -->|feeds| {nid(tname)}")
    for s, targets in sorted(ext_srcs.items()):
        for t in sorted(targets):
            lines.append(f"  {nid(s)} -->|loads| {nid(t)}")
    lines += ["```", ""]
    (out_root / "ER_DIAGRAM_LINEAGE.md").write_text("\n".join(lines), encoding="utf-8")


def generate_layer(layer_id: str, design: dict, tpl_dir: Path, out_root: Path) -> str:
    layer = design["layers"][layer_id]
    out_dir = out_root / layer_id
    out_dir.mkdir(parents=True, exist_ok=True)
    all_rels = relationships(design)
    layer_rels = [r for r in all_rels if r["layer"] == layer_id]
    gen_sttm(layer_id, layer, tpl_dir, out_dir)
    gen_dm(layer_id, layer, tpl_dir, out_dir)
    gen_dq(layer_id, layer, tpl_dir, out_dir)
    gen_er_md(layer_id, layer, layer_rels, out_dir)
    gen_er_xlsx(design, all_rels, tpl_dir, out_root)
    gen_lineage_md(design, out_root)
    cols = sum(len(t["columns"]) for t in layer["tables"].values())
    rules = sum(len(t["rules"]) for t in layer["tables"].values())
    return (f"OK layer={layer_id} tables={len(layer['tables'])} columns={cols} "
            f"rules={rules} rels={len(layer_rels)}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--layer", required=True,
                    help="layer id from column_mappings.json (e.g. L0), or 'all'")
    ap.add_argument("--design-dir", required=True, type=Path,
                    help="outputs/00_state/design (column_mappings.json + dq_rules_design.json)")
    ap.add_argument("--out-root", required=True, type=Path,
                    help="outputs/05_artifacts (per-layer folders are created under it)")
    ap.add_argument("--templates-dir", required=True, type=Path,
                    help=".claude/skills/design-agent/templates/workbooks")
    args = ap.parse_args()

    if not args.templates_dir.exists():
        raise SystemExit(f"FATAL: templates dir not found: {args.templates_dir}")
    args.out_root.mkdir(parents=True, exist_ok=True)

    design = load_design(args.design_dir)
    if args.layer.lower() == "all":
        layer_ids = list(design["layers"])
    else:
        if args.layer not in design["layers"]:
            raise SystemExit(f"FATAL: layer '{args.layer}' not in design (have: {list(design['layers'])})")
        layer_ids = [args.layer]

    for layer_id in layer_ids:
        print(generate_layer(layer_id, design, args.templates_dir, args.out_root))
    return 0


if __name__ == "__main__":
    sys.exit(main())
