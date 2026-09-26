#!/usr/bin/env python3
"""Verify generated workbook headers match canonical templates exactly.

Compares row 1 of every sheet in every generated ``*.xlsx`` under ``--artifacts-dir``
against the matching canonical template under ``--templates-dir``. Any missing
sheet or header mismatch is a fatal error and returned via exit code 1 so the
orchestrator can halt the phase.

Usage:
    python verify_workbook_headers.py --artifacts-dir outputs/05_artifacts
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - environment dependent
    load_workbook = None  # type: ignore[assignment]


TEMPLATE_BY_PREFIX = {
    "STTM": "STTM_TEMPLATE.xlsx",
    "DATA_MODEL": "DATA_MODEL_TEMPLATE.xlsx",
    "DQ": "DQ_TEMPLATE.xlsx",
    "ER_DIAGRAM": "ER_DIAGRAM_TEMPLATE.xlsx",
}


def classify_workbook(path: Path) -> str | None:
    name = path.name.upper()
    for prefix in sorted(TEMPLATE_BY_PREFIX, key=len, reverse=True):
        if name.startswith(prefix):
            return prefix
    return None


def row_values(sheet: Any) -> list[Any]:
    return [cell.value for cell in sheet[1]]


def compare_workbook(workbook_path: Path, template_path: Path) -> list[dict[str, Any]]:
    generated = load_workbook(workbook_path, read_only=True, data_only=False)
    template = load_workbook(template_path, read_only=True, data_only=False)
    issues: list[dict[str, Any]] = []

    for sheet_name in template.sheetnames:
        if sheet_name.startswith("_"):
            # _template_reference and similar underscore sheets are internal
            # scaffold sheets that the generator renames/deletes before save.
            # Their absence in the generated workbook is correct behaviour.
            continue
        if sheet_name not in generated.sheetnames:
            issues.append(
                {
                    "workbook": str(workbook_path),
                    "sheet": sheet_name,
                    "issue": "missing_sheet",
                }
            )
            continue

        generated_header = row_values(generated[sheet_name])
        template_header = row_values(template[sheet_name])
        max_len = max(len(generated_header), len(template_header))
        generated_header += [None] * (max_len - len(generated_header))
        template_header += [None] * (max_len - len(template_header))

        for index, (actual, expected) in enumerate(zip(generated_header, template_header), start=1):
            if actual != expected:
                issues.append(
                    {
                        "workbook": str(workbook_path),
                        "sheet": sheet_name,
                        "column_index": index,
                        "expected": expected,
                        "actual": actual,
                    }
                )

    generated.close()
    template.close()
    return issues


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--artifacts-dir",
        required=True,
        help="Directory containing generated .xlsx artifacts (recursively scanned)",
    )
    parser.add_argument(
        "--templates-dir",
        default=".claude/skills/design-agent/templates/workbooks",
        help="Directory containing canonical *_TEMPLATE.xlsx files",
    )
    parser.add_argument("--report", help="Optional JSON report path")
    args = parser.parse_args()

    if load_workbook is None:
        print("openpyxl is required to verify workbook headers", file=sys.stderr)
        return 3

    artifacts_dir = Path(args.artifacts_dir)
    templates_dir = Path(args.templates_dir)
    report_path = Path(args.report) if args.report else None

    issues: list[dict[str, Any]] = []
    checked: list[str] = []
    if not artifacts_dir.exists():
        print(f"Artifacts directory not found: {artifacts_dir}", file=sys.stderr)
        return 2

    for workbook_path in sorted(artifacts_dir.rglob("*.xlsx")):
        workbook_type = classify_workbook(workbook_path)
        if workbook_type is None:
            continue
        template_path = templates_dir / TEMPLATE_BY_PREFIX[workbook_type]
        if not template_path.exists():
            issues.append(
                {
                    "workbook": str(workbook_path),
                    "template": str(template_path),
                    "issue": "missing_template",
                }
            )
            continue
        checked.append(str(workbook_path))
        issues.extend(compare_workbook(workbook_path, template_path))

    payload = {"checked": checked, "issues": issues}
    if report_path:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    if issues:
        print(json.dumps(payload, indent=2), file=sys.stderr)
        return 1

    print(json.dumps(payload, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
